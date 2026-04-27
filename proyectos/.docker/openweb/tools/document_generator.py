"""
title: Document Generator
author: shimmer
description: Genera documentos Word (.docx), Excel (.xlsx), CSV, TXT, MD, JSON y XML descargables desde el chat.
requirements: python-docx, openpyxl
version: 0.8.0
"""

import base64
import io
import csv
import json
import xml.dom.minidom
from docx import Document
from fastapi.responses import HTMLResponse
from openpyxl import Workbook
from pydantic import BaseModel


def _download_iframe(b64: str, filename: str, mime: str) -> HTMLResponse:
    ext = filename.rsplit(".", 1)[-1].upper()
    size_kb = round(len(base64.b64decode(b64)) / 1024, 1)

    icons = {
        "DOCX": "📄", "XLSX": "📊", "CSV": "📋",
        "TXT": "📝", "MD": "📝", "JSON": "🔧", "XML": "🔧",
    }
    colors = {
        "DOCX": "#2B579A", "XLSX": "#217346", "CSV": "#6B7280",
        "TXT": "#374151", "MD": "#374151", "JSON": "#B45309", "XML": "#7C3AED",
    }
    icon = icons.get(ext, "📁")
    color = colors.get(ext, "#4F46E5")

    html = f"""<!DOCTYPE html>
<html>
<head>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
    background: transparent;
    padding: 4px;
  }}
  .card {{
    display: flex;
    align-items: center;
    gap: 14px;
    background: #ffffff;
    border: 1px solid #e5e7eb;
    border-radius: 12px;
    padding: 14px 16px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.07);
  }}
  .icon {{ font-size: 2rem; flex-shrink: 0; }}
  .info {{ flex: 1; min-width: 0; }}
  .filename {{
    font-size: 0.92rem;
    font-weight: 600;
    color: #111827;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
  }}
  .meta {{ font-size: 0.75rem; color: #6b7280; margin-top: 2px; }}
  .btn {{
    flex-shrink: 0;
    background: {color};
    color: white;
    border: none;
    border-radius: 8px;
    padding: 8px 16px;
    font-size: 0.82rem;
    font-weight: 600;
    cursor: pointer;
    transition: opacity 0.15s;
    white-space: nowrap;
  }}
  .btn:hover {{ opacity: 0.85; }}
</style>
</head>
<body>
<div class="card">
  <div class="icon">{icon}</div>
  <div class="info">
    <div class="filename">{filename}</div>
    <div class="meta">{ext} · {size_kb} KB · Listo para descargar</div>
  </div>
  <button class="btn" onclick="download()">⬇ Descargar</button>
</div>
<script>
function download() {{
  const b64 = "{b64}";
  const bin = atob(b64);
  const arr = new Uint8Array(bin.length);
  for (let i = 0; i < bin.length; i++) arr[i] = bin.charCodeAt(i);
  const blob = new Blob([arr], {{type: "{mime}"}});
  const url = URL.createObjectURL(blob);
  const a = document.createElement("a");
  a.href = url;
  a.download = "{filename}";
  document.body.appendChild(a);
  a.click();
  document.body.removeChild(a);
  URL.revokeObjectURL(url);
}}
</script>
<script>
  function reportHeight() {{
    parent.postMessage({{ type: 'iframe:height', height: document.documentElement.scrollHeight }}, '*');
  }}
  window.addEventListener('load', reportHeight);
  new ResizeObserver(reportHeight).observe(document.body);
</script>
</body>
</html>"""
    return HTMLResponse(content=html, headers={"Content-Disposition": "inline"})


class Tools:
    class Valves(BaseModel):
        pass

    def __init__(self):
        self.valves = self.Valves()

    def generate_docx(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Word (.docx) con el contenido proporcionado.
        :param content: Texto o markdown (# título, ## sección, - lista).
        :param filename: Nombre del archivo sin extensión.
        """
        doc = Document()
        for line in content.split("\n"):
            line = line.strip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line:
                doc.add_paragraph(line)

        buffer = io.BytesIO()
        doc.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".docx",
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    def generate_excel(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Excel (.xlsx) a partir de datos en formato CSV.
        La primera fila se trata como cabecera en negrita.
        :param data: Datos CSV. Ejemplo: "Nombre,Edad\nAlice,30\nBob,25"
        :param filename: Nombre del archivo sin extensión.
        """
        wb = Workbook()
        ws = wb.active
        for i, row in enumerate(csv.reader(io.StringIO(data))):
            ws.append(row)
            if i == 0:
                for cell in ws[1]:
                    cell.font = cell.font.copy(bold=True)

        buffer = io.BytesIO()
        wb.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".xlsx",
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def generate_csv(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo CSV descargable.
        :param data: Datos CSV. Ejemplo: "Nombre,Edad\nAlice,30\nBob,25"
        :param filename: Nombre del archivo sin extensión.
        """
        b64 = base64.b64encode(data.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".csv", "text/csv;charset=utf-8")

    def generate_txt(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo de texto plano (.txt) descargable.
        :param content: Contenido de texto plano.
        :param filename: Nombre del archivo sin extensión.
        """
        b64 = base64.b64encode(content.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".txt", "text/plain;charset=utf-8")

    def generate_md(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Markdown (.md) descargable.
        :param content: Contenido en formato Markdown.
        :param filename: Nombre del archivo sin extensión.
        """
        b64 = base64.b64encode(content.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".md", "text/markdown;charset=utf-8")

    def generate_json(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo JSON (.json) descargable con formato indentado.
        :param data: Contenido JSON como string. Se reindentará automáticamente.
        :param filename: Nombre del archivo sin extensión.
        """
        try:
            parsed = json.loads(data)
            formatted = json.dumps(parsed, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            formatted = data

        b64 = base64.b64encode(formatted.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".json", "application/json;charset=utf-8")

    def generate_xml(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo XML (.xml) descargable con formato indentado.
        :param data: Contenido XML como string. Se reindentará automáticamente.
        :param filename: Nombre del archivo sin extensión.
        """
        try:
            formatted = xml.dom.minidom.parseString(data).toprettyxml(indent="  ")
        except Exception:
            formatted = data

        b64 = base64.b64encode(formatted.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".xml", "application/xml;charset=utf-8")
