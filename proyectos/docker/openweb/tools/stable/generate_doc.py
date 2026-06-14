"""
title: Document Generator
author: shimmer
description: Genera documentos Word (.docx), Excel (.xlsx), CSV, MD, JSON y XML descargables desde el chat.
requirements: python-docx, openpyxl
version: 0.9.0
"""

import base64
import io
import csv
import json
import re
import xml.dom.minidom
from docx import Document
from fastapi.responses import HTMLResponse
from openpyxl import Workbook
from pydantic import BaseModel


def _add_inline_runs(para, text: str):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            para.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            para.add_run(token[1:-1]).italic = True
        elif token:
            para.add_run(token)


def _add_paragraph_with_inline(doc, text: str, style: str = None):
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    _add_inline_runs(para, text)
    return para


def _download_iframe(b64: str, filename: str, mime: str) -> HTMLResponse:
    parts = filename.rsplit(".", 1)
    ext = parts[-1].upper() if len(parts) > 1 else "FILE"
    size_kb = round(len(base64.b64decode(b64)) / 1024, 1)
    # dos variantes del nombre: una para el html y otra para el atributo js — evita xss y rotura de string
    filename_html = filename.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    filename_js = filename.replace("\\", "\\\\").replace('"', '\\"')

    icons = {
        "DOCX": "📄", "XLSX": "📊", "CSV": "📋",
        "MD": "📝", "JSON": "🔧", "XML": "🔧",
    }
    colors = {
        "DOCX": "#2B579A", "XLSX": "#217346", "CSV": "#6B7280",
        "MD": "#374151", "JSON": "#B45309", "XML": "#7C3AED",
    }
    icon = icons.get(ext, "📁")
    color = colors.get(ext, "#4F46E5")

    html = f"""<!DOCTYPE html>
<!-- tarjeta de descarga renderizada como iframe en el chat de open webui -->
<html>
<head>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
    background: transparent;
    padding: 4px;
  }}
  .card {{
    display: flex;
    align-items: center;
    gap: 14px;
    width: fit-content;
    background: #ffffff;
    border: 1px solid #e5e7eb;
    border-radius: 12px;
    padding: 14px 16px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.07);
  }}
  .icon {{ font-size: 2rem; flex-shrink: 0; }}
  .info {{ flex: 1; min-width: 0; }}
  .filename {{
    font-size: 0.92rem;
    font-weight: 600;
    color: #111827;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
  }}
  .meta {{ font-size: 0.75rem; color: #6b7280; margin-top: 2px; }}
  .btn {{
    flex-shrink: 0;
    background: {color};
    color: white;
    border: none;
    border-radius: 8px;
    padding: 8px 16px;
    font-size: 0.82rem;
    font-weight: 600;
    cursor: pointer;
    transition: opacity 0.15s;
    white-space: nowrap;
  }}
  .btn:hover {{ opacity: 0.85; }}
</style>
</head>
<body>
<div class="card">
  <div class="icon">{icon}</div>
  <button class="btn" onclick="download()">⬇ Descargar</button>
  <div class="info">
    <div class="filename">{filename_html}</div>
    <div class="meta">{ext} · {size_kb} KB</div>
  </div>
</div>
<script>
function download() {{
  const b64 = "{b64}";
  const bin = atob(b64);
  const arr = new Uint8Array(bin.length);
  for (let i = 0; i < bin.length; i++) arr[i] = bin.charCodeAt(i);
  const blob = new Blob([arr], {{type: "{mime}"}});
  const url = URL.createObjectURL(blob);
  const a = document.createElement("a");
  a.href = url;
  a.download = "{filename_js}";
  document.body.appendChild(a);
  a.click();
  document.body.removeChild(a);
  URL.revokeObjectURL(url);
}}
</script>
<script>
  function reportHeight() {{
    parent.postMessage({{ type: 'iframe:height', height: document.documentElement.scrollHeight }}, '*');
  }}
  window.addEventListener('load', reportHeight);
  new ResizeObserver(reportHeight).observe(document.body);
</script>
</body>
</html>"""
    return HTMLResponse(content=html, headers={"Content-Disposition": "inline"})


class Tools:
    class Valves(BaseModel):
        pass

    def __init__(self):
        self.valves = self.Valves()

    def generate_docx(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Word (.docx) con el contenido proporcionado.
        :param content: Texto o markdown (# título, ## sección, - lista).
        :param filename: Nombre del archivo sin extensión.
        """
        content = content.encode("utf-16", errors="surrogatepass").decode("utf-16")
        filename = filename.encode("utf-16", errors="surrogatepass").decode("utf-16")
        doc = Document()
        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            if line.startswith("|") and line.endswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                # excluir la fila separadora |---|---| que markdown requiere pero word no
                rows = [r for r in table_lines if not re.match(r'^\|[\s\-:| ]+\|$', r)]
                if rows:
                    cols = [c.strip() for c in rows[0].strip("|").split("|")]
                    table = doc.add_table(rows=len(rows), cols=len(cols))
                    table.style = "Table Grid"
                    for r_idx, row_line in enumerate(rows):
                        cells = [c.strip() for c in row_line.strip("|").split("|")]
                        for c_idx, cell_text in enumerate(cells):
                            if c_idx < len(table.rows[r_idx].cells):
                                cell = table.rows[r_idx].cells[c_idx]
                                cell.paragraphs[0].clear()
                                _add_inline_runs(cell.paragraphs[0], cell_text)
                                if r_idx == 0:
                                    for run in cell.paragraphs[0].runs:
                                        run.bold = True
                continue

            if re.match(r'^-{3,}$', line):
                i += 1
                continue
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("> "):
                para = doc.add_paragraph(style="Quote") if "Quote" in [s.name for s in doc.styles] else doc.add_paragraph()
                _add_inline_runs(para, line[2:])
            elif line.startswith("- ") or line.startswith("* "):
                _add_paragraph_with_inline(doc, line[2:], style="List Bullet")
            elif re.match(r'^\d+\.\s', line):
                _add_paragraph_with_inline(doc, re.sub(r'^\d+\.\s', '', line), style="List Number")
            elif line:
                _add_paragraph_with_inline(doc, line)
            i += 1

        buffer = io.BytesIO()
        doc.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".docx",
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    def generate_excel(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Excel (.xlsx) a partir de datos en formato CSV.
        La primera fila se trata como cabecera en negrita.
        Usa ; como separador de columnas para evitar conflictos con comas en el contenido.
        :param data: Datos CSV con ; como separador. Los campos que contengan ; deben ir entre comillas dobles.
                     Ejemplo: "Nombre;Skills\nAlice;\"Python;Java\"\nBob;SQL"
        :param filename: Nombre del archivo sin extensión.
        """
        wb = Workbook()
        ws = wb.active
        # ; como separador porque las celdas pueden contener comas (listas, precios, nombres)
        for i, row in enumerate(csv.reader(io.StringIO(data), delimiter=";")):
            ws.append(row)
            if i == 0:
                for cell in ws[1]:
                    cell.font = cell.font.copy(bold=True)

        buffer = io.BytesIO()
        wb.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".xlsx",
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def generate_csv(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo CSV descargable.
        Usa ; como separador de columnas para evitar conflictos con comas en el contenido.
        :param data: Datos CSV con ; como separador. Ejemplo: "Nombre;Edad\nAlice;30\nBob;25"
        :param filename: Nombre del archivo sin extensión.
        """
        b64 = base64.b64encode(data.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".csv", "text/csv;charset=utf-8")

    def generate_md(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Markdown (.md) descargable.
        Para contenido HTML: NO uses esta tool. Escribe el HTML directamente en el chat.
        :param content: Contenido en formato Markdown.
        :param filename: Nombre del archivo sin extensión.
        """
        b64 = base64.b64encode(content.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".md", "text/markdown;charset=utf-8")

    def generate_json(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo JSON (.json) descargable con formato indentado.
        :param data: Contenido JSON como string. Se reindentará automáticamente.
        :param filename: Nombre del archivo sin extensión.
        """
        try:
            parsed = json.loads(data)
            formatted = json.dumps(parsed, ensure_ascii=False, indent=2)
        except json.JSONDecodeError as e:
            formatted = f"// JSON inválido: {e}\n{data}"

        b64 = base64.b64encode(formatted.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".json", "application/json;charset=utf-8")

    def generate_xml(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo XML (.xml) descargable con formato indentado.
        :param data: Contenido XML como string. Se reindentará automáticamente.
        :param filename: Nombre del archivo sin extensión.
        """
        try:
            formatted = xml.dom.minidom.parseString(data).toprettyxml(indent="  ")
        except Exception:
            formatted = data

        b64 = base64.b64encode(formatted.encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, filename.replace(" ", "_") + ".xml", "application/xml;charset=utf-8")
