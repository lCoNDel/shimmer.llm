"""
title: Universal Document Generator
author: shimmer.llm
description: Genera cualquier tipo de documento descargable desde el chat — Word, Excel (multi-hoja, con tipos numéricos reales), PowerPoint (con notas de orador), PDF (con bloques de código y citas), CSV, Markdown, JSON, XML, HTML y texto/código. Diseñada para ser usada por agentes IA.
requirements: python-docx, openpyxl, python-pptx, fpdf2
version: 1.1.0
"""

import base64
import csv
import io
import json
import re
import xml.dom.minidom

from docx import Document
from docx.shared import Pt
from fastapi.responses import HTMLResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel


# ── Utilidades comunes ─────────────────────────────────────────────────────────

def _clean_filename(filename: str) -> str:
    filename = filename.encode("utf-16", errors="surrogatepass").decode("utf-16")
    filename = re.sub(r'[\\/:*?"<>|]', "", filename).strip()
    return (filename or "documento").replace(" ", "_")


def _clean_text(text: str) -> str:
    return text.encode("utf-16", errors="surrogatepass").decode("utf-16")


def _strip_md(text: str) -> str:
    """Elimina marcadores inline de Markdown (** * `) en contextos que no los renderizan."""
    return re.sub(r"\*\*?|`", "", text)


_NUM_RE_EN = re.compile(r"^-?\d+(\.\d+)?$")
_NUM_RE_ES = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$|^-?\d+,\d+$")


def _coerce_number(value: str):
    """Convierte strings numéricos (formato inglés o español) a int/float para que Excel pueda operar.
    Conserva como texto los códigos con ceros a la izquierda (ej. '007')."""
    s = value.strip()
    if not s:
        return value
    body = s.lstrip("-")
    if body.startswith("0") and len(body) > 1 and body[1] not in ".,":
        return value
    if _NUM_RE_EN.match(s):
        return float(s) if "." in s else int(s)
    if _NUM_RE_ES.match(s):
        return float(s.replace(".", "").replace(",", "."))
    return value


def _download_iframe(b64: str, filename: str, mime: str) -> HTMLResponse:
    parts = filename.rsplit(".", 1)
    ext = parts[-1].upper() if len(parts) > 1 else "FILE"
    size_kb = round(len(base64.b64decode(b64)) / 1024, 1)
    # dos variantes del nombre: una para el html y otra para el atributo js — evita xss y rotura de string
    filename_html = filename.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    filename_js = filename.replace("\\", "\\\\").replace('"', '\\"')

    icons = {
        "DOCX": "📄", "XLSX": "📊", "CSV": "📋", "PPTX": "📽️", "PDF": "📕",
        "MD": "📝", "JSON": "🔧", "XML": "🔧", "HTML": "🌐", "TXT": "📃",
    }
    colors = {
        "DOCX": "#2B579A", "XLSX": "#217346", "CSV": "#6B7280", "PPTX": "#D24726",
        "PDF": "#B91C1C", "MD": "#374151", "JSON": "#B45309", "XML": "#7C3AED",
        "HTML": "#0E7490", "TXT": "#4B5563",
    }
    icon = icons.get(ext, "📁")
    color = colors.get(ext, "#4F46E5")

    html = f"""<!DOCTYPE html>
<!-- tarjeta de descarga renderizada como iframe en el chat de open webui -->
<html>
<head>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
    background: transparent;
    padding: 4px;
  }}
  .card {{
    display: flex;
    align-items: center;
    gap: 14px;
    width: fit-content;
    background: #ffffff;
    border: 1px solid #e5e7eb;
    border-radius: 12px;
    padding: 14px 16px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.07);
  }}
  .icon {{ font-size: 2rem; flex-shrink: 0; }}
  .info {{ flex: 1; min-width: 0; }}
  .filename {{
    font-size: 0.92rem;
    font-weight: 600;
    color: #111827;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
  }}
  .meta {{ font-size: 0.75rem; color: #6b7280; margin-top: 2px; }}
  .btn {{
    flex-shrink: 0;
    background: {color};
    color: white;
    border: none;
    border-radius: 8px;
    padding: 8px 16px;
    font-size: 0.82rem;
    font-weight: 600;
    cursor: pointer;
    transition: opacity 0.15s;
    white-space: nowrap;
  }}
  .btn:hover {{ opacity: 0.85; }}
</style>
</head>
<body>
<div class="card">
  <div class="icon">{icon}</div>
  <button class="btn" onclick="download()">⬇ Descargar</button>
  <div class="info">
    <div class="filename">{filename_html}</div>
    <div class="meta">{ext} · {size_kb} KB</div>
  </div>
</div>
<script>
function download() {{
  const b64 = "{b64}";
  const bin = atob(b64);
  const arr = new Uint8Array(bin.length);
  for (let i = 0; i < bin.length; i++) arr[i] = bin.charCodeAt(i);
  const blob = new Blob([arr], {{type: "{mime}"}});
  const url = URL.createObjectURL(blob);
  const a = document.createElement("a");
  a.href = url;
  a.download = "{filename_js}";
  document.body.appendChild(a);
  a.click();
  document.body.removeChild(a);
  URL.revokeObjectURL(url);
}}
</script>
<script>
  function reportHeight() {{
    parent.postMessage({{ type: 'iframe:height', height: document.documentElement.scrollHeight }}, '*');
  }}
  window.addEventListener('load', reportHeight);
  new ResizeObserver(reportHeight).observe(document.body);
</script>
</body>
</html>"""
    return HTMLResponse(content=html, headers={"Content-Disposition": "inline"})


def _text_response(content: str, filename: str, ext: str, mime: str) -> HTMLResponse:
    b64 = base64.b64encode(content.encode("utf-8")).decode("utf-8")
    return _download_iframe(b64, f"{filename}.{ext}", mime)


# ── Markdown → Word ────────────────────────────────────────────────────────────

def _add_inline_runs(para, text: str):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            para.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = para.add_run(token[1:-1])
            run.font.name = "Courier New"
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            para.add_run(token[1:-1]).italic = True
        elif token:
            para.add_run(token)


def _add_paragraph_with_inline(doc, text: str, style: str = None):
    try:
        para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    except KeyError:
        para = doc.add_paragraph()
    _add_inline_runs(para, text)
    return para


# ── Markdown → PDF ─────────────────────────────────────────────────────────────

_PDF_REPLACEMENTS = {
    "€": "EUR", "“": '"', "”": '"', "‘": "'", "’": "'", "–": "-", "—": "-",
    "…": "...", "•": "-", "→": "->", "←": "<-", "≥": ">=", "≤": "<=", "≠": "!=",
    "✓": "[OK]", "✗": "[X]", "º": "o", "ª": "a",
}


def _pdf_safe(text: str) -> str:
    for k, v in _PDF_REPLACEMENTS.items():
        text = text.replace(k, v)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_markdown(text: str) -> str:
    # fpdf2 usa **bold** y __italic__; markdown estándar usa *italic*
    text = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"__\1__", text)
    return text


# ── Tool ───────────────────────────────────────────────────────────────────────

class Tools:
    class Valves(BaseModel):
        pass

    def __init__(self):
        self.valves = self.Valves()

    # ── Word ───────────────────────────────────────────────────────────────────

    def generate_docx(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Word (.docx) a partir de Markdown.
        Soporta: # ## ### #### títulos, **negrita**, *cursiva*, `código`, - listas,
        1. listas numeradas, > citas, tablas |a|b|, bloques ``` de código y
        [PAGEBREAK] para salto de página.
        :param content: Contenido en Markdown.
        :param filename: Nombre del archivo sin extensión.
        """
        content = _clean_text(content)
        filename = _clean_filename(filename)
        doc = Document()
        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            if line.startswith("```"):
                i += 1
                code_lines = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                i += 1  # cierre ```
                for code_line in code_lines or [""]:
                    para = doc.add_paragraph()
                    run = para.add_run(code_line if code_line else " ")
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)
                    para.paragraph_format.space_after = Pt(0)
                continue

            if line.startswith("|") and line.endswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                rows = [r for r in table_lines if not re.match(r'^\|[\s\-:| ]+\|$', r)]
                if rows:
                    cols = [c.strip() for c in rows[0].strip("|").split("|")]
                    table = doc.add_table(rows=len(rows), cols=len(cols))
                    table.style = "Table Grid"
                    for r_idx, row_line in enumerate(rows):
                        cells = [c.strip() for c in row_line.strip("|").split("|")]
                        for c_idx, cell_text in enumerate(cells):
                            if c_idx < len(table.rows[r_idx].cells):
                                cell = table.rows[r_idx].cells[c_idx]
                                cell.paragraphs[0].clear()
                                _add_inline_runs(cell.paragraphs[0], cell_text)
                                if r_idx == 0:
                                    for run in cell.paragraphs[0].runs:
                                        run.bold = True
                continue

            if line == "[PAGEBREAK]":
                doc.add_page_break()
            elif re.match(r'^-{3,}$', line):
                pass
            elif line.startswith("#### "):
                doc.add_heading(_strip_md(line[5:]), level=4)
            elif line.startswith("### "):
                doc.add_heading(_strip_md(line[4:]), level=3)
            elif line.startswith("## "):
                doc.add_heading(_strip_md(line[3:]), level=2)
            elif line.startswith("# "):
                doc.add_heading(_strip_md(line[2:]), level=1)
            elif line.startswith("> "):
                _add_paragraph_with_inline(doc, line[2:], style="Quote")
            elif re.match(r'^\s*[-*]\s', lines[i]):
                indent = len(lines[i]) - len(lines[i].lstrip())
                style = "List Bullet 2" if indent >= 2 else "List Bullet"
                _add_paragraph_with_inline(doc, re.sub(r'^\s*[-*]\s', '', lines[i]), style=style)
            elif re.match(r'^\d+\.\s', line):
                _add_paragraph_with_inline(doc, re.sub(r'^\d+\.\s', '', line), style="List Number")
            elif line:
                _add_paragraph_with_inline(doc, line)
            i += 1

        buffer = io.BytesIO()
        doc.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename + ".docx",
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    # ── PDF ────────────────────────────────────────────────────────────────────

    def generate_pdf(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo PDF a partir de Markdown.
        Soporta: # ## ### títulos, **negrita**, *cursiva*, - listas, 1. listas numeradas,
        tablas |a|b|, > citas, bloques ``` de código, --- como línea separadora
        y [PAGEBREAK] para salto de página.
        :param content: Contenido en Markdown.
        :param filename: Nombre del archivo sin extensión.
        """
        from fpdf import FPDF, XPos, YPos

        content = _clean_text(content)
        filename = _clean_filename(filename)

        pdf = FPDF(format="A4")
        pdf.set_auto_page_break(auto=True, margin=18)
        pdf.set_margins(18, 18, 18)
        pdf.add_page()
        pdf.set_font("helvetica", size=11)

        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            if line.startswith("```"):
                i += 1
                code_lines = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                i += 1  # cierre ```
                pdf.set_font("courier", size=8.5)
                for code_line in code_lines or [" "]:
                    pdf.multi_cell(0, 4.5, _pdf_safe(code_line) or " ", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font("helvetica", size=11)
                pdf.ln(2)
                continue

            if line.startswith("|") and line.endswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                rows = [r for r in table_lines if not re.match(r'^\|[\s\-:| ]+\|$', r)]
                if rows:
                    data = [
                        [_pdf_safe(_strip_md(c.strip()))
                         for c in r.strip("|").split("|")]
                        for r in rows
                    ]
                    pdf.set_font("helvetica", size=9)
                    with pdf.table() as table:
                        for row_data in data:
                            row = table.row()
                            for cell in row_data:
                                row.cell(cell)
                    pdf.set_font("helvetica", size=11)
                    pdf.ln(3)
                continue

            if line == "[PAGEBREAK]":
                pdf.add_page()
            elif re.match(r'^-{3,}$', line):
                pdf.ln(2)
                pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
                pdf.ln(4)
            elif line.startswith("# "):
                pdf.set_font("helvetica", "B", 17)
                pdf.multi_cell(0, 9, _pdf_safe(_strip_md(line[2:])), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font("helvetica", size=11)
                pdf.ln(2)
            elif line.startswith("## "):
                pdf.set_font("helvetica", "B", 14)
                pdf.multi_cell(0, 8, _pdf_safe(_strip_md(line[3:])), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font("helvetica", size=11)
                pdf.ln(1)
            elif line.startswith("### ") or line.startswith("#### "):
                pdf.set_font("helvetica", "B", 12)
                pdf.multi_cell(0, 7, _pdf_safe(_strip_md(line.lstrip("#").strip())), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font("helvetica", size=11)
            elif line.startswith("> "):
                pdf.set_font("helvetica", "I", 11)
                pdf.set_x(pdf.l_margin + 6)
                pdf.multi_cell(0, 6, _pdf_safe(_strip_md(line[2:])), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font("helvetica", size=11)
            elif re.match(r'^[-*]\s', line):
                pdf.multi_cell(0, 6, _pdf_safe(_pdf_markdown("   - " + line[2:])), markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            elif re.match(r'^\d+\.\s', line):
                pdf.multi_cell(0, 6, _pdf_safe(_pdf_markdown("   " + line)), markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            elif line:
                pdf.multi_cell(0, 6, _pdf_safe(_pdf_markdown(line)), markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            else:
                pdf.ln(3)
            i += 1

        b64 = base64.b64encode(bytes(pdf.output())).decode("utf-8")
        return _download_iframe(b64, filename + ".pdf", "application/pdf")

    # ── PowerPoint ─────────────────────────────────────────────────────────────

    def generate_pptx(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera una presentación PowerPoint (.pptx).
        Formato: cada "# Título" inicia una diapositiva nueva; las líneas "- punto"
        debajo son sus viñetas (dos espacios de sangría = sub-viñeta). La primera
        diapositiva puede llevar "## subtítulo" para crear una portada. Las líneas
        "> texto" se añaden como notas del orador de la diapositiva actual.
        Ejemplo: "# Portada\\n## Subtítulo\\n# Tema 1\\n- punto A\\n> nota del orador"
        :param content: Contenido estructurado en Markdown.
        :param filename: Nombre del archivo sin extensión.
        """
        from pptx import Presentation
        from pptx.util import Inches, Pt

        content = _clean_text(content)
        filename = _clean_filename(filename)

        # agrupar en diapositivas
        slides = []
        current = None
        for raw in content.split("\n"):
            line = raw.rstrip()
            stripped = line.strip()
            if stripped.startswith("# "):
                current = {"title": _strip_md(stripped[2:]), "subtitle": "", "bullets": [], "notes": ""}
                slides.append(current)
            elif current is None:
                continue
            elif stripped.startswith("## "):
                current["subtitle"] = _strip_md(stripped[3:])
            elif stripped.startswith("> "):
                current["notes"] = (current["notes"] + "\n" + _strip_md(stripped[2:])).strip()
            elif re.match(r'^\s*[-*]\s', line):
                indent = len(line) - len(line.lstrip())
                text = re.sub(r'^\s*[-*]\s', '', line)
                current["bullets"].append((1 if indent >= 2 else 0, text))
            elif stripped:
                current["bullets"].append((0, stripped))

        if not slides:
            slides = [{"title": "Documento", "subtitle": "", "bullets": [(0, l) for l in content.split("\n") if l.strip()], "notes": ""}]

        prs = Presentation()
        for idx, s in enumerate(slides):
            if idx == 0 and s["subtitle"] and not s["bullets"]:
                layout = prs.slide_layouts[0]  # portada
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = s["title"]
                slide.placeholders[1].text = s["subtitle"]
                if s["notes"]:
                    slide.notes_slide.notes_text_frame.text = s["notes"]
                continue
            layout = prs.slide_layouts[1]  # título + contenido
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s["title"]
            body = slide.placeholders[1].text_frame
            items = s["bullets"] or ([(0, s["subtitle"])] if s["subtitle"] else [])
            first = True
            for level, text in items:
                para = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                para.level = level
                para.text = _strip_md(text)
                para.font.size = Pt(20 if level == 0 else 16)
            if s["notes"]:
                slide.notes_slide.notes_text_frame.text = s["notes"]

        buffer = io.BytesIO()
        prs.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename + ".pptx",
                                "application/vnd.openxmlformats-officedocument.presentationml.presentation")

    # ── Excel ──────────────────────────────────────────────────────────────────

    def generate_excel(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Excel (.xlsx). Acepta dos formatos de entrada:
        1) CSV con ; como separador (una sola hoja). Ejemplo: "Nombre;Edad\\nAna;30"
        2) JSON multi-hoja: {"Hoja1": [["col1","col2"],["a",1]], "Hoja2": [...]}.
           Las filas también pueden ser objetos: {"Hoja1": [{"col1":"a","col2":1}]}.
        Cabeceras en negrita, panel congelado, autofiltro y ancho de columna automático.
        En modo CSV los valores numéricos ("1.234,56" o "1234.56") se escriben como
        números reales para que Excel pueda sumar y filtrar.
        :param data: Datos en CSV (separador ;) o JSON multi-hoja.
        :param filename: Nombre del archivo sin extensión.
        """
        data = _clean_text(data)
        filename = _clean_filename(filename)

        sheets = {}
        stripped = data.strip()
        if stripped.startswith("{"):
            try:
                parsed = json.loads(stripped)
                if isinstance(parsed, dict):
                    for sheet_name, rows in parsed.items():
                        if not isinstance(rows, list):
                            continue
                        if rows and isinstance(rows[0], dict):
                            headers = list(rows[0].keys())
                            sheets[str(sheet_name)[:31]] = [headers] + [
                                [r.get(h, "") for h in headers] for r in rows if isinstance(r, dict)
                            ]
                        else:
                            sheets[str(sheet_name)[:31]] = [
                                r if isinstance(r, list) else [r] for r in rows
                            ]
            except json.JSONDecodeError:
                pass

        if not sheets:
            raw_rows = list(csv.reader(io.StringIO(data), delimiter=";"))
            rows = [raw_rows[0]] if raw_rows else []
            rows += [[_coerce_number(c) for c in row] for row in raw_rows[1:]]
            sheets = {"Hoja1": rows}

        wb = Workbook()
        wb.remove(wb.active)
        for sheet_name, rows in sheets.items():
            ws = wb.create_sheet(title=sheet_name or "Hoja")
            widths = {}
            for row in rows:
                ws.append(row)
                for c_idx, value in enumerate(row, start=1):
                    width = len(str(value)) if value is not None else 0
                    widths[c_idx] = min(max(widths.get(c_idx, 8), width + 2), 60)
            if rows:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                ws.freeze_panes = "A2"
                if len(rows) > 1:
                    last_col = get_column_letter(max(len(r) for r in rows if r) or 1)
                    ws.auto_filter.ref = f"A1:{last_col}{len(rows)}"
            for c_idx, width in widths.items():
                ws.column_dimensions[get_column_letter(c_idx)].width = width

        if not wb.sheetnames:
            wb.create_sheet("Hoja1")

        buffer = io.BytesIO()
        wb.save(buffer)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        return _download_iframe(b64, filename + ".xlsx",
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ── CSV ────────────────────────────────────────────────────────────────────

    def generate_csv(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo CSV descargable (separador ; — formato español/Excel).
        Los campos que contengan ; deben ir entre comillas dobles.
        :param data: Datos CSV con ; como separador. Ejemplo: "Nombre;Edad\\nAna;30"
        :param filename: Nombre del archivo sin extensión.
        """
        data = _clean_text(data)
        # BOM para que Excel abra el UTF-8 correctamente
        b64 = base64.b64encode(("\ufeff" + data).encode("utf-8")).decode("utf-8")
        return _download_iframe(b64, _clean_filename(filename) + ".csv", "text/csv;charset=utf-8")

    # ── Markdown ───────────────────────────────────────────────────────────────

    def generate_md(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo Markdown (.md) descargable.
        :param content: Contenido en formato Markdown.
        :param filename: Nombre del archivo sin extensión.
        """
        return _text_response(_clean_text(content), _clean_filename(filename), "md", "text/markdown;charset=utf-8")

    # ── JSON ───────────────────────────────────────────────────────────────────

    def generate_json(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo JSON (.json) descargable con formato indentado.
        :param data: Contenido JSON como string. Se reindentará automáticamente.
        :param filename: Nombre del archivo sin extensión.
        """
        data = _clean_text(data)
        try:
            formatted = json.dumps(json.loads(data), ensure_ascii=False, indent=2)
        except json.JSONDecodeError as e:
            formatted = f"// JSON inválido: {e}\n{data}"
        return _text_response(formatted, _clean_filename(filename), "json", "application/json;charset=utf-8")

    # ── XML ────────────────────────────────────────────────────────────────────

    def generate_xml(self, data: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo XML (.xml) descargable con formato indentado.
        :param data: Contenido XML como string. Se reindentará automáticamente.
        :param filename: Nombre del archivo sin extensión.
        """
        data = _clean_text(data)
        try:
            formatted = xml.dom.minidom.parseString(data).toprettyxml(indent="  ")
        except Exception:
            formatted = data
        return _text_response(formatted, _clean_filename(filename), "xml", "application/xml;charset=utf-8")

    # ── HTML ───────────────────────────────────────────────────────────────────

    def generate_html(self, content: str, filename: str = "documento") -> HTMLResponse:
        """
        Genera un archivo HTML (.html) descargable. Si el contenido no incluye
        estructura <html>, lo envuelve en una plantilla básica con estilos limpios.
        :param content: Contenido HTML (completo o solo el body).
        :param filename: Nombre del archivo sin extensión.
        """
        content = _clean_text(content)
        if "<html" not in content.lower():
            content = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
  body {{ font-family: -apple-system, 'Segoe UI', sans-serif; max-width: 860px; margin: 2rem auto; padding: 0 1rem; line-height: 1.6; color: #1f2937; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #d1d5db; padding: 6px 10px; text-align: left; }}
  th {{ background: #f3f4f6; }}
  code {{ background: #f3f4f6; padding: 2px 5px; border-radius: 4px; }}
</style>
</head>
<body>
{content}
</body>
</html>"""
        return _text_response(content, _clean_filename(filename), "html", "text/html;charset=utf-8")

    # ── Texto / código genérico ────────────────────────────────────────────────

    def generate_text_file(self, content: str, filename: str = "documento", extension: str = "txt") -> HTMLResponse:
        """
        Genera un archivo de texto plano o código con cualquier extensión:
        txt, log, py, js, sql, yaml, ini, cfg, ps1, sh, bat, css, etc.
        Úsala para cualquier formato que no tenga una función específica.
        :param content: Contenido del archivo tal cual.
        :param filename: Nombre del archivo sin extensión.
        :param extension: Extensión sin punto (por defecto "txt").
        """
        content = _clean_text(content)
        ext = re.sub(r'[^a-z0-9]', '', extension.lower().lstrip(".")) or "txt"
        mimes = {
            "txt": "text/plain", "log": "text/plain", "py": "text/x-python",
            "js": "text/javascript", "css": "text/css", "sql": "application/sql",
            "yaml": "application/yaml", "yml": "application/yaml",
            "ini": "text/plain", "cfg": "text/plain", "ps1": "text/plain",
            "sh": "text/x-shellscript", "bat": "text/plain",
        }
        mime = mimes.get(ext, "text/plain") + ";charset=utf-8"
        return _text_response(content, _clean_filename(filename), ext, mime)
